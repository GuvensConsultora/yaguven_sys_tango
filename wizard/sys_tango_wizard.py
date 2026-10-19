import base64
import io

from odoo import api, fields, models
from odoo.exceptions import UserError

# Grupos internos de Odoo que se presentan como cuentas patrimoniales.
PATRIMONIAL_GROUPS = ("asset", "liability", "equity", "off")

COLUMNS = [
    ("code", "Código de cuenta"),
    ("name", "Descripción de cuenta"),
    ("debit", "Débitos acumulados"),
    ("credit", "Créditos acumulados"),
    ("balance_debit", "Saldo deudor"),
    ("balance_credit", "Saldo acreedor"),
    ("asset", "Activo"),
    ("liability", "Pasivo"),
    ("loss", "Resultado negativo"),
    ("profit", "Resultado positivo"),
]
AMOUNT_KEYS = [k for k, _ in COLUMNS[2:]]


class SysTangoWizard(models.TransientModel):
    _name = "yaguven.sys.tango.wizard"
    _description = "Sumas y saldos con plan Tango"

    company_id = fields.Many2one(
        "res.company",
        string="Compañía",
        required=True,
        default=lambda self: self.env.company,
    )
    date_from = fields.Date(
        string="Desde",
        help="Inicio del ejercicio. Lo anterior entra como saldo inicial de cada "
        "cuenta patrimonial; los resultados de ejercicios anteriores van a "
        "Resultados no asignados, igual que en el balance nativo de Odoo.",
    )
    date_to = fields.Date(
        string="Hasta",
        required=True,
        default=fields.Date.context_today,
    )
    show_detail = fields.Boolean(
        string="Mostrar cuentas Odoo",
        help="Debajo de cada cuenta Tango, las cuentas Odoo que la componen.",
    )
    show_zero = fields.Boolean(
        string="Incluir cuentas sin saldo",
        default=True,
        help="Lista todas las cuentas del plan Tango, como el SyS de Tango.",
    )
    line_ids = fields.One2many(
        "yaguven.sys.tango.line", "wizard_id", string="Renglones"
    )

    @api.onchange("company_id", "date_to")
    def _onchange_dates(self):
        if self.company_id and self.date_to:
            self.date_from = self.company_id.compute_fiscalyear_dates(
                self.date_to
            )["date_from"]

    @api.model_create_multi
    def create(self, vals_list):
        # Sin «Desde» explícito: inicio del ejercicio de la fecha «Hasta» pedida.
        for vals in vals_list:
            if not vals.get("date_from"):
                company = self.env["res.company"].browse(
                    vals.get("company_id") or self.env.company.id
                )
                date_to = fields.Date.to_date(
                    vals.get("date_to")
                ) or fields.Date.context_today(self)
                vals["date_from"] = company.compute_fiscalyear_dates(date_to)[
                    "date_from"
                ]
        return super().create(vals_list)

    # ------------------------------------------------------------------ cálculo

    def _excluded_moves(self):
        """Asientos que no entran al reporte: los que tocan una cuenta tildada
        «Excluir del SyS Tango», con todas sus contrapartidas. Devuelve también
        cuánto quedó afuera, para informarlo al pie."""
        self.ensure_one()
        accounts = self.env["account.account"].search([
            ("company_ids", "in", self.company_id.id),
            ("sys_tango_excluded", "=", True),
        ])
        info = {
            "accounts": accounts, "move_ids": [], "moves": 0,
            "count": 0, "debit": 0.0, "prior_count": 0,
        }
        if not accounts:
            return info
        AML = self.env["account.move.line"]
        base = [
            ("company_id", "=", self.company_id.id),
            ("parent_state", "=", "posted"),
            ("date", "<=", self.date_to),
        ]
        # Los posteriores a «Hasta» no entran al reporte: no hace falta excluirlos.
        move_ids = AML.search(base + [("account_id", "in", accounts.ids)]).move_id.ids
        if not move_ids:
            return info
        info["move_ids"], info["moves"] = move_ids, len(move_ids)
        # El importe se informa sobre la misma base que la columna «Débitos
        # acumulados»: sólo las líneas del período. Las anteriores al «Desde»
        # entran al listado neteadas en el saldo inicial, no como débitos, así
        # que se cuentan aparte y no se suman a ese importe.
        groups = AML._read_group(
            base + [
                ("move_id", "in", move_ids),
                ("date", ">=", self.date_from),
            ],
            [],
            ["debit:sum", "__count"],
        )
        debit, count = groups[0] if groups else (0.0, 0)
        prior = AML.search_count(
            base + [("move_id", "in", move_ids), ("date", "<", self.date_from)]
        )
        info.update(count=count, debit=debit, prior_count=prior)
        return info

    def _account_amounts(self, split_accounts, excluded_move_ids=()):
        """{(account, journal): [debit, credit]} al estilo del SyS: el saldo
        inicial neto entra como un débito o un crédito, más los movimientos del
        período. Sólo las cuentas de split_accounts (las que tienen equivalencia
        por diario) se separan por diario; el resto va con journal vacío, así
        el neteo del saldo inicial queda igual que por cuenta."""
        self.ensure_one()
        no_journal = self.env["account.journal"]
        if self.date_from > self.date_to:
            raise UserError("La fecha Desde es posterior a la fecha Hasta.")
        AML = self.env["account.move.line"].with_company(self.company_id)
        base = [
            ("company_id", "=", self.company_id.id),
            ("parent_state", "=", "posted"),
        ]
        if excluded_move_ids:
            base.append(("move_id", "not in", list(excluded_move_ids)))
        amounts = {}

        def add(account, journal, debit, credit):
            key = (account, journal if account in split_accounts else no_journal)
            acc = amounts.setdefault(key, [0.0, 0.0])
            acc[0] += debit
            acc[1] += credit

        # Saldo inicial: patrimoniales con toda su historia; resultados de
        # ejercicios anteriores acumulados en Resultados no asignados.
        # Se netea por clave (cuenta, o cuenta+diario si se reparte).
        prior_pnl = 0.0
        initial = {}
        for account, journal, debit, credit in AML._read_group(
            base + [("date", "<", self.date_from)],
            ["account_id", "journal_id"],
            ["debit:sum", "credit:sum"],
        ):
            if not account.include_initial_balance:
                prior_pnl += debit - credit
                continue
            key = (account, journal if account in split_accounts else no_journal)
            initial[key] = initial.get(key, 0.0) + debit - credit
        for (account, journal), net in initial.items():
            add(account, journal, max(net, 0.0), max(-net, 0.0))
        if not self.company_id.currency_id.is_zero(prior_pnl):
            unaffected = self.env["account.account"].with_company(
                self.company_id
            ).search(
                [
                    ("company_ids", "in", self.company_id.id),
                    ("account_type", "=", "equity_unaffected"),
                ],
                limit=1,
            )
            if not unaffected:
                raise UserError(
                    "Hay resultados de ejercicios anteriores y la compañía no "
                    "tiene cuenta de Resultados no asignados."
                )
            add(unaffected, no_journal, max(prior_pnl, 0.0), max(-prior_pnl, 0.0))

        for account, journal, debit, credit in AML._read_group(
            base + [
                ("date", ">=", self.date_from),
                ("date", "<=", self.date_to),
            ],
            ["account_id", "journal_id"],
            ["debit:sum", "credit:sum"],
        ):
            add(account, journal, debit, credit)
        return amounts

    @staticmethod
    def _amount_row(debit, credit, nature):
        balance = debit - credit
        row = {
            "debit": debit,
            "credit": credit,
            "balance_debit": max(balance, 0.0),
            "balance_credit": max(-balance, 0.0),
            "asset": 0.0,
            "liability": 0.0,
            "loss": 0.0,
            "profit": 0.0,
        }
        if nature == "patrimonial":
            row["asset"], row["liability"] = row["balance_debit"], row["balance_credit"]
        else:
            row["loss"], row["profit"] = row["balance_debit"], row["balance_credit"]
        return row

    def _compute_rows(self, exclusion=None):
        """Renglones del reporte, en el orden del plan Tango.

        kind: 'tango' cuenta Tango · 'detail' cuenta Odoo que la compone ·
        'unmapped' cuenta Odoo sin equivalencia · 'total'.
        """
        self.ensure_one()
        exclusion = exclusion if exclusion is not None else self._excluded_moves()
        company = self.company_id
        rounding = company.currency_id
        mappings = self.env["yaguven.tango.mapping"].search(
            [("company_id", "=", company.id)]
        )
        tango_of = {
            (m.account_id, m.journal_id): m.tango_account_id for m in mappings
        }
        split_accounts = {m.account_id for m in mappings if m.journal_id}
        amounts = self._account_amounts(split_accounts, exclusion["move_ids"])
        no_journal = self.env["account.journal"]
        tango_accounts = self.env["yaguven.tango.account"].search(
            [("company_id", "=", company.id)]
        )
        if not tango_accounts:
            raise UserError(
                f"La compañía {company.name} no tiene cargado el plan Tango."
            )

        # Cada importe va a la equivalencia de su diario y, si no hay, a la
        # general de la cuenta. Las partes de una misma cuenta que caen en la
        # equivalencia general se juntan en un solo renglón de detalle.
        by_tango = {t: {} for t in tango_accounts}
        unmapped = {}
        for (account, journal), (debit, credit) in amounts.items():
            if rounding.is_zero(debit) and rounding.is_zero(credit):
                continue
            tango = journal and tango_of.get((account, journal))
            part_key = (account, journal) if tango else (account, no_journal)
            if not tango:
                tango = tango_of.get((account, no_journal))
            target = by_tango[tango] if tango in by_tango else unmapped
            part = target.setdefault(part_key, [0.0, 0.0])
            part[0] += debit
            part[1] += credit

        def detail_label(account, journal):
            name = account.name
            if journal:
                name = f"{name} — {journal.name}"
            return account.code, name

        rows = []
        totals = dict.fromkeys(AMOUNT_KEYS, 0.0)
        for tango in tango_accounts:
            parts = sorted(
                ((acc, jou, d, c) for (acc, jou), (d, c) in by_tango[tango].items()),
                key=lambda p: (p[0].code or "", p[1].name or ""),
            )
            debit = sum(p[2] for p in parts)
            credit = sum(p[3] for p in parts)
            if not parts and not self.show_zero:
                continue
            row = self._amount_row(debit, credit, tango.nature)
            rows.append(dict(row, kind="tango", code=tango.code, name=tango.name))
            for key in AMOUNT_KEYS:
                totals[key] += row[key]
            if self.show_detail:
                for account, journal, p_debit, p_credit in parts:
                    code, name = detail_label(account, journal)
                    rows.append(dict(
                        self._amount_row(p_debit, p_credit, tango.nature),
                        kind="detail",
                        code=code,
                        name=name,
                    ))

        # Sin equivalencia siempre se agrupa por cuenta (journal vacío).
        for (account, _journal), (debit, credit) in sorted(
            unmapped.items(), key=lambda p: p[0][0].code or ""
        ):
            nature = (
                "patrimonial"
                if account.internal_group in PATRIMONIAL_GROUPS
                else "resultado"
            )
            row = self._amount_row(debit, credit, nature)
            rows.append(dict(
                row,
                kind="unmapped",
                code=account.code,
                name=f"SIN EQUIVALENCIA TANGO — {account.name}",
            ))
            for key in AMOUNT_KEYS:
                totals[key] += row[key]

        rows.append(dict(totals, kind="total", code="", name="TOTALES"))
        return rows

    @staticmethod
    def _fmt(amount):
        """Importe con el formato de la plaza: 1.234.567,89."""
        return f"{amount:,.2f}".replace(",", "X").replace(".", ",").replace("X", ".")

    def _check_totals(self, rows, exclusion=None):
        """Controles de cuadre del reporte, como texto para mostrar."""
        total = next(r for r in rows if r["kind"] == "total")
        currency = self.company_id.currency_id
        checks = [
            ("Débitos = Créditos", total["debit"] - total["credit"]),
            (
                "Saldo deudor = Saldo acreedor",
                total["balance_debit"] - total["balance_credit"],
            ),
            (
                "Activo − Pasivo = Resultado positivo − Resultado negativo",
                (total["asset"] - total["liability"])
                - (total["profit"] - total["loss"]),
            ),
        ]
        unmapped = sum(1 for r in rows if r["kind"] == "unmapped")
        lines = [
            f"{name}: {'OK' if currency.is_zero(diff) else f'DIFERENCIA {self._fmt(diff)}'}"
            for name, diff in checks
        ]
        lines.append(
            "Cuentas Odoo sin equivalencia Tango: "
            + (f"{unmapped} (ver renglones «SIN EQUIVALENCIA»)" if unmapped else "ninguna")
        )
        exclusion = exclusion if exclusion is not None else self._excluded_moves()
        if exclusion["accounts"]:
            codes = ", ".join(
                f"{a.code} {a.name}" for a in exclusion["accounts"][:6]
            ) + (" …" if len(exclusion["accounts"]) > 6 else "")
            lines.append(
                f"Cuentas excluidas del SyS Tango: {len(exclusion['accounts'])} ({codes})"
            )
            lines.append(
                f"Asientos que las tocan y quedaron afuera: "
                f"{exclusion['moves']} asientos · {exclusion['count']} líneas del "
                f"período · débitos {self._fmt(exclusion['debit'])}"
            )
            if exclusion["prior_count"]:
                lines.append(
                    f"De esos asientos hay además {exclusion['prior_count']} líneas "
                    "anteriores al «Desde», que tampoco entran al saldo inicial."
                )
            lines.append(
                "Por esas exclusiones este listado NO coincide con el balance de "
                "sumas y saldos de Odoo: se quitaron los asientos completos, "
                "incluidas las contrapartidas en otras cuentas."
            )
        else:
            lines.append("Cuentas excluidas del SyS Tango: ninguna")
        return lines

    # ------------------------------------------------------------------ salidas

    def action_show(self):
        self.ensure_one()
        exclusion = self._excluded_moves()
        rows = self._compute_rows(exclusion)
        # En pantalla los controles van como renglones, debajo de TOTALES.
        rows += [
            dict(dict.fromkeys(AMOUNT_KEYS, 0.0), kind="info", code="", name=line)
            for line in self._check_totals(rows, exclusion)
        ]
        self.line_ids.unlink()
        self.env["yaguven.sys.tango.line"].create([
            dict(
                {k: r[k] for k in AMOUNT_KEYS},
                wizard_id=self.id,
                sequence=idx,
                kind=r["kind"],
                code=r["code"],
                name=r["name"],
            )
            for idx, r in enumerate(rows)
        ])
        return {
            "type": "ir.actions.act_window",
            "name": f"Sumas y saldos Tango — {self.company_id.name} "
            f"{self.date_from.strftime('%d/%m/%Y')} a {self.date_to.strftime('%d/%m/%Y')}",
            "res_model": "yaguven.sys.tango.line",
            "view_mode": "list",
            "domain": [("wizard_id", "=", self.id)],
            "context": {"create": False, "edit": False, "delete": False},
            "target": "current",
        }

    def action_export_xlsx(self):
        self.ensure_one()
        exclusion = self._excluded_moves()
        rows = self._compute_rows(exclusion)

        import openpyxl
        from openpyxl.styles import Alignment, Font, PatternFill

        wb = openpyxl.Workbook()
        ws = wb.active
        ws.title = "Balance de saldos"

        header_fill = PatternFill("solid", fgColor="6C7A92")
        header_font = Font(bold=True, color="FFFFFF")
        bold = Font(bold=True)
        detail_font = Font(italic=True, color="555555")
        unmapped_font = Font(bold=True, color="C65A00")
        num_fmt = "#,##0.00"

        company = self.company_id
        ws.append([f"{company.name} — CUIT {company.vat or ''}"])
        ws.append([
            "Balance de sumas y saldos con plan de cuentas Tango — "
            f"del {self.date_from.strftime('%d/%m/%Y')} "
            f"al {self.date_to.strftime('%d/%m/%Y')}"
        ])
        ws.append([])
        ws.cell(1, 1).font = bold
        header_row = ws.max_row + 1
        ws.append([label for _, label in COLUMNS])
        for cell in ws[header_row]:
            cell.fill = header_fill
            cell.font = header_font
            cell.alignment = Alignment(horizontal="center", wrap_text=True)

        for r in rows:
            if r["kind"] == "info":  # en el xlsx van abajo, en «Controles»
                continue
            code = r["code"]
            if r["kind"] == "detail":
                code = f"    {code}"
            ws.append([code, r["name"]] + [r[k] for k in AMOUNT_KEYS])
            current = ws.max_row
            for col in range(3, len(COLUMNS) + 1):
                ws.cell(current, col).number_format = num_fmt
            font = {
                "detail": detail_font,
                "unmapped": unmapped_font,
                "total": bold,
            }.get(r["kind"])
            if font:
                for cell in ws[current]:
                    cell.font = font

        ws.append([])
        ws.append(["Controles"])
        ws.cell(ws.max_row, 1).font = bold
        for line in self._check_totals(rows, exclusion):
            ws.append([line])

        ws.column_dimensions["A"].width = 18
        ws.column_dimensions["B"].width = 48
        for letter in "CDEFGHIJ":
            ws.column_dimensions[letter].width = 18
        ws.freeze_panes = ws.cell(header_row + 1, 3)

        buf = io.BytesIO()
        wb.save(buf)
        buf.seek(0)
        filename = (
            f"SyS_Tango_{company.name}_{self.date_to.strftime('%Y%m%d')}.xlsx"
        ).replace(" ", "_")
        attachment = self.env["ir.attachment"].create({
            "name": filename,
            "type": "binary",
            "datas": base64.b64encode(buf.read()),
            "mimetype": "application/vnd.openxmlformats-officedocument."
            "spreadsheetml.sheet",
        })
        return {
            "type": "ir.actions.act_url",
            "url": f"/web/content/{attachment.id}?download=true",
            "target": "self",
        }


class SysTangoLine(models.TransientModel):
    _name = "yaguven.sys.tango.line"
    _description = "Renglón de sumas y saldos con plan Tango"
    _order = "sequence"

    wizard_id = fields.Many2one(
        "yaguven.sys.tango.wizard", required=True, ondelete="cascade"
    )
    sequence = fields.Integer()
    kind = fields.Selection(
        [
            ("tango", "Cuenta Tango"),
            ("detail", "Cuenta Odoo"),
            ("unmapped", "Sin equivalencia"),
            ("total", "Totales"),
            ("info", "Control"),
        ],
        string="Renglón",
    )
    code = fields.Char(string="Código de cuenta")
    name = fields.Char(string="Descripción de cuenta")
    currency_id = fields.Many2one(related="wizard_id.company_id.currency_id")
    debit = fields.Monetary(string="Débitos acumulados")
    credit = fields.Monetary(string="Créditos acumulados")
    balance_debit = fields.Monetary(string="Saldo deudor")
    balance_credit = fields.Monetary(string="Saldo acreedor")
    asset = fields.Monetary(string="Activo")
    liability = fields.Monetary(string="Pasivo")
    loss = fields.Monetary(string="Resultado negativo")
    profit = fields.Monetary(string="Resultado positivo")
